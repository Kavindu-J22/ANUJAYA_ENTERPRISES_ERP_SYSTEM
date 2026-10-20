import openpyxl
import json

wb_val = openpyxl.load_workbook('China.xlsx', data_only=True)
wb_formula = openpyxl.load_workbook('China.xlsx', data_only=False)

output = {}

for name in wb_val.sheetnames:
    sheet_val = wb_val[name]
    sheet_form = wb_formula[name]
    rows = []
    for r in range(1, sheet_val.max_row + 1):
        row_vals = [sheet_val.cell(r, c).value for c in range(1, sheet_val.max_column + 1)]
        row_forms = [sheet_form.cell(r, c).value for c in range(1, sheet_form.max_column + 1)]
        if any(v is not None for v in row_vals):
            rows.append({
                "row": r,
                "values": [str(v) if v is not None else None for v in row_vals],
                "formulas": [str(f) if f is not None else None for f in row_forms]
            })
    output[name] = rows

with open('scratch_china_parsed.json', 'w', encoding='utf-8') as f:
    json.dump(output, f, indent=2, ensure_ascii=False)

print("Exported scratch_china_parsed.json successfully!")
